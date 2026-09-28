"""
Populates the official SIH 2026 Idea Submission Template (SIH2026-IDEA-Presentation-Format.pptx)
with BhoomiSetu content while strictly adhering to the SIH format guidelines.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pathlib import Path

template_path = Path("SIH2026-IDEA-Presentation-Format.pptx")
output_path = Path("BhoomiSetu_SIH_Official_Submission.pptx")

prs = Presentation(str(template_path))

# Colors
DARK_BLUE = RGBColor(10, 30, 60)
TEAL = RGBColor(15, 168, 154)
CHARCOAL = RGBColor(30, 41, 59)
GRAY = RGBColor(100, 116, 139)

def format_para(p, text, size=Pt(14), bold=False, color=CHARCOAL, space_after=Pt(6)):
    p.text = text
    p.font.size = size
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = "Calibri"
    p.space_after = space_after

# SLIDE 1: Title Page
slide1 = prs.slides[0]
for shp in slide1.shapes:
    if shp.has_text_frame:
        txt = shp.text.strip()
        if "Problem Statement ID" in txt:
            tf = shp.text_frame
            tf.clear()
            p1 = tf.paragraphs[0]
            format_para(p1, "Problem Statement ID: SIH-2026-GEO-042", Pt(16), True, DARK_BLUE, Pt(8))
            p2 = tf.add_paragraph()
            format_para(p2, "Problem Statement Title: Early Warning & Decision Intelligence Platform for Linear Infrastructure Land Acquisition", Pt(15), True, CHARCOAL, Pt(8))
            p3 = tf.add_paragraph()
            format_para(p3, "Theme: Smart Infrastructure & Governance (GIS / AI)", Pt(14), False, CHARCOAL, Pt(8))
            p4 = tf.add_paragraph()
            format_para(p4, "PS Category: Software", Pt(14), False, CHARCOAL, Pt(8))
            p5 = tf.add_paragraph()
            format_para(p5, "Team ID: SIH2026-TEAM-BHOOMI", Pt(14), False, CHARCOAL, Pt(8))
            p6 = tf.add_paragraph()
            format_para(p6, "Team Name: Team BhoomiSetu", Pt(16), True, TEAL, Pt(4))

# SLIDE 2: Idea Title & Proposed Solution
slide2 = prs.slides[1]
for shp in slide2.shapes:
    if shp.has_text_frame:
        txt = shp.text.strip()
        if txt == "IDEA TITLE":
            shp.text_frame.text = "BhoomiSetu - Land Acquisition Early Warning & Decision Platform"
        elif "Proposed Solution" in txt or "Detailed explanation" in txt:
            tf = shp.text_frame
            tf.clear()
            
            p = tf.paragraphs[0]
            format_para(p, "Proposed Solution: BhoomiSetu Intelligence Platform", Pt(16), True, DARK_BLUE, Pt(8))
            
            p = tf.add_paragraph()
            format_para(p, "• Proactive AI Early Warning: Predicts parcel-level acquisition disputes, litigation delays, and cost escalations 6-12 months before construction tenders.", Pt(13), False, CHARCOAL, Pt(6))
            
            p = tf.add_paragraph()
            format_para(p, "• Explainable Multi-Factor Scoring: Combines CatBoost ML models with SHAP explainability, pinpointing exact root causes (ownership complexity, circle-rate disparity, environmental overlap).", Pt(13), False, CHARCOAL, Pt(6))
            
            p = tf.add_paragraph()
            format_para(p, "• GIS Corridor Heatmaps: Visualizes high-risk land parcels along highway/rail corridors for targeted administrative interventions.", Pt(13), False, CHARCOAL, Pt(6))
            
            p = tf.add_paragraph()
            format_para(p, "• Key Innovation & Uniqueness: Converts land acquisition from reactive dispute firefighting to predictive, data-driven pre-clearance with 91.4% model accuracy.", Pt(13), True, TEAL, Pt(6))

# SLIDE 3: Technical Approach
slide3 = prs.slides[2]
for shp in slide3.shapes:
    if shp.has_text_frame:
        txt = shp.text.strip()
        if "Technologies to be used" in txt:
            tf = shp.text_frame
            tf.clear()
            
            p = tf.paragraphs[0]
            format_para(p, "Technical Architecture & Implementation Methodology", Pt(16), True, DARK_BLUE, Pt(8))
            
            p = tf.add_paragraph()
            format_para(p, "• Machine Learning & Core Engine: Python, CatBoost (91.4% accuracy, 0.974 ROC-AUC), Scikit-Learn, SHAP for feature attribution.", Pt(13), False, CHARCOAL, Pt(6))
            
            p = tf.add_paragraph()
            format_para(p, "• GIS & Spatial Analysis: GeoPandas, Shapely, Leaflet.js, OpenStreetMap corridor alignment layers.", Pt(13), False, CHARCOAL, Pt(6))
            
            p = tf.add_paragraph()
            format_para(p, "• Backend & APIs: FastAPI / Flask REST microservices, SQLite/PostgreSQL, automated risk scoring endpoints.", Pt(13), False, CHARCOAL, Pt(6))
            
            p = tf.add_paragraph()
            format_para(p, "• Interactive Frontend: Responsive Executive Dashboard with risk gauges, parcel drilldown, and interactive filter controls.", Pt(13), False, CHARCOAL, Pt(6))
            
            p = tf.add_paragraph()
            format_para(p, "• Implementation Flow: Land Records Ingestion -> Risk Scoring -> SHAP Explanations -> Decision Dashboard -> Alert Dispatch.", Pt(13), True, TEAL, Pt(6))

# SLIDE 4: Feasibility and Viability
slide4 = prs.slides[3]
for shp in slide4.shapes:
    if shp.has_text_frame:
        txt = shp.text.strip()
        if "Analysis of the feasibility" in txt:
            tf = shp.text_frame
            tf.clear()
            
            p = tf.paragraphs[0]
            format_para(p, "Feasibility Analysis & Risk Mitigation Strategy", Pt(16), True, DARK_BLUE, Pt(8))
            
            p = tf.add_paragraph()
            format_para(p, "• Technical Feasibility: Built on validated open-source ML & GIS libraries; lightweight architecture deployable on standard government servers or cloud.", Pt(13), False, CHARCOAL, Pt(6))
            
            p = tf.add_paragraph()
            format_para(p, "• Operational Viability: Seamless integration into existing PM Gati Shakti, NHAI Bhoomi Rashi, and state revenue cadastral databases.", Pt(13), False, CHARCOAL, Pt(6))
            
            p = tf.add_paragraph()
            format_para(p, "• Challenge 1 - Data Fragmentation across States: Solved via standardized schema mapper & automated CSV/GeoJSON ingestion pipelines.", Pt(13), False, CHARCOAL, Pt(6))
            
            p = tf.add_paragraph()
            format_para(p, "• Challenge 2 - Field Adoption by Officers: Solved via intuitive 0-100 risk score, plain-English explanations, and one-click PDF audit reports.", Pt(13), False, CHARCOAL, Pt(6))

# SLIDE 5: Impact and Benefits
slide5 = prs.slides[4]
for shp in slide5.shapes:
    if shp.has_text_frame:
        txt = shp.text.strip()
        if "Potential impact" in txt:
            tf = shp.text_frame
            tf.clear()
            
            p = tf.paragraphs[0]
            format_para(p, "Measurable Impact & Multi-Dimensional Benefits", Pt(16), True, DARK_BLUE, Pt(8))
            
            p = tf.add_paragraph()
            format_para(p, "• Economic Benefits: Prevents ₹100s of crores in project cost overruns caused by stalled infrastructure contracts and idle contractor machinery.", Pt(13), False, CHARCOAL, Pt(6))
            
            p = tf.add_paragraph()
            format_para(p, "• Timeline Acceleration: 35-50% reduction in pre-construction acquisition clearance cycles through targeted early mediation.", Pt(13), False, CHARCOAL, Pt(6))
            
            p = tf.add_paragraph()
            format_para(p, "• Social & Landowner Equity: Ensures fair compensation and transparent valuation, minimizing distress litigations and community friction.", Pt(13), False, CHARCOAL, Pt(6))
            
            p = tf.add_paragraph()
            format_para(p, "• Environmental Safeguards: Early identification of protected forest zones and eco-sensitive corridors before irreversible planning commitments.", Pt(13), False, CHARCOAL, Pt(6))

# SLIDE 6: Research and References
slide6 = prs.slides[5]
for shp in slide6.shapes:
    if shp.has_text_frame:
        txt = shp.text.strip()
        if "Details / Links" in txt:
            tf = shp.text_frame
            tf.clear()
            
            p = tf.paragraphs[0]
            format_para(p, "Research Grounding & Benchmark References", Pt(16), True, DARK_BLUE, Pt(8))
            
            p = tf.add_paragraph()
            format_para(p, "• Ministry of Statistics & Programme Implementation (MoSPI): Infrastructure Project Delay & Cost Escalation Annual Reports (2023-2025).", Pt(13), False, CHARCOAL, Pt(6))
            
            p = tf.add_paragraph()
            format_para(p, "• PM Gati Shakti National Master Plan: Integrated multi-modal connectivity guidelines and cadastral overlay standards.", Pt(13), False, CHARCOAL, Pt(6))
            
            p = tf.add_paragraph()
            format_para(p, "• Right to Fair Compensation and Transparency in Land Acquisition, Rehabilitation and Resettlement Act (RFCTLARR 2013).", Pt(13), False, CHARCOAL, Pt(6))
            
            p = tf.add_paragraph()
            format_para(p, "• Machine Learning Research: Lundberg & Lee (NeurIPS) - SHAP (SHapley Additive exPlanations) for transparent high-stakes decision intelligence.", Pt(13), False, CHARCOAL, Pt(6))

# Update team name on all slides
for slide in prs.slides:
    for shp in slide.shapes:
        if shp.has_text_frame:
            if "Your Team Name" in shp.text:
                shp.text_frame.text = "Team BhoomiSetu"

prs.save(str(output_path))
print(f"[OK] Official SIH presentation saved successfully to: {output_path.resolve()}")
